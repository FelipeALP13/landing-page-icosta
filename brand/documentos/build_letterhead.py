from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image


ROOT = Path(__file__).resolve().parents[2]
BRAND = ROOT / "brand"
OUTPUT = Path(__file__).resolve().parent / "folha-timbrada-icosta.docx"

GREEN = "061A15"
INK = "10221C"
GOLD = "D5B66C"


def set_run_font(run, name, size, color, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.font.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)


def add_field(paragraph, instruction, size=6.5, bold=False):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction_node, separate, value, end])
    set_run_font(run, "Arial", size, INK, bold=bold)


def clear_container(container):
    element = container._element
    for child in list(element):
        element.remove(child)


def set_picture_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", title)
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def logo_for_light_background():
    image = Image.open(BRAND / "logo-primary.png").convert("RGBA")
    pixels = []
    for red, green, blue, alpha in image.get_flattened_data():
        # The primary logo uses warm white for the initial and descriptor.
        # Recolor only neutral light pixels so the mark is legible on white.
        if alpha and max(red, green, blue) - min(red, green, blue) < 32 and max(red, green, blue) > 90:
            pixels.append((14, 59, 48, alpha))
        else:
            pixels.append((red, green, blue, alpha))
    image.putdata(pixels)
    stream = BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    return stream


def set_paragraph_top_border(paragraph, color, size=6, space=3):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    top = borders.find(qn("w:top"))
    if top is None:
        top = OxmlElement("w:top")
        borders.append(top)
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), str(space))
    top.set(qn("w:color"), color)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.orientation = 0
    section.top_margin = Mm(45)
    section.bottom_margin = Mm(24)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(5)

    core = document.core_properties
    core.title = "Folha timbrada ICosta Advocacia"
    core.subject = "Papel timbrado institucional"
    core.author = "ICosta Advocacia e Consultoria Jurídica"
    core.keywords = "ICosta; papel timbrado; advocacia; consultoria jurídica"

    normal = document.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(11.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.15

    header = section.header
    clear_container(header)
    logo_paragraph = header.add_paragraph()
    logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    logo_paragraph.paragraph_format.left_indent = Mm(-10)
    logo_paragraph.paragraph_format.right_indent = Mm(-10)
    logo_paragraph.paragraph_format.space_after = Pt(0)
    shape = logo_paragraph.add_run().add_picture(logo_for_light_background(), width=Mm(102))
    set_picture_alt_text(shape, "Logotipo ICosta", "ICosta Advocacia e Consultoria Jurídica")

    footer = section.footer
    clear_container(footer)
    page = footer.add_paragraph()
    page.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page.paragraph_format.left_indent = Mm(-15)
    page.paragraph_format.right_indent = Mm(-15)
    page.paragraph_format.space_after = Pt(1)
    set_run_font(page.add_run("página "), "Arial", 6.5, INK)
    add_field(page, "PAGE")
    set_run_font(page.add_run(" de "), "Arial", 6.5, INK)
    add_field(page, "NUMPAGES")

    info = footer.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.paragraph_format.left_indent = Mm(-15)
    info.paragraph_format.right_indent = Mm(-15)
    info.paragraph_format.space_before = Pt(0)
    info.paragraph_format.space_after = Pt(0)
    set_paragraph_top_border(info, GOLD, size=6, space=3)
    set_run_font(
        info.add_run("Av. Alfredo Ignácio Nogueira Penido, 335, sala 104, Jd. Aquarius, São José dos Campos/SP, CEP 12246-000"),
        "Arial",
        6.5,
        INK,
    )
    set_run_font(info.add_run("  |  "), "Arial", 6.5, GOLD)
    set_run_font(info.add_run("(12) 3922-9758 / 98103-9999"), "Arial", 6.5, INK)
    set_run_font(info.add_run("  |  "), "Arial", 6.5, GOLD)
    set_run_font(info.add_run("contato@icosta.adv.br"), "Arial", 6.5, INK)

    body = document.add_paragraph()
    body.paragraph_format.space_after = Pt(0)

    document.settings.element.append(OxmlElement("w:updateFields"))
    document.settings.element[-1].set(qn("w:val"), "true")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
